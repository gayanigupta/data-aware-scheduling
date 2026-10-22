#!/usr/bin/env python3
"""Build revisions/MILP_solver_details.docx -- the task-2 write-up.
Run after milp_solver_report.py has produced the CSV:
    python3 revisions/make_task2_report.py
"""

import os

import pandas as pd
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(__file__)
CSV = os.path.join(HERE, "results", "task-02", "milp_solver_details.csv")

df = pd.read_csv(CSV)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

doc.add_heading("MILP solver details for every run", 0)
doc.add_paragraph(
    "For every MILP run in the paper's experiments, we recorded the size "
    "of the model handed to the solver (number of variables and "
    "constraints), what the solver reported when it stopped, how long it "
    "ran, and the optimality gap it had left when it stopped. All rows "
    "are in one file: revisions/results/task-02/milp_solver_details.csv.")

doc.add_heading("What is in the file", 1)
for s in [
    "experiment - which study the run belongs to (deadline slack, "
    "federation size, workflow size, task failures, link bandwidth, or a "
    "per-dataset test run),",
    "param - the setting on that study's axis (e.g. alpha = 1.2, or "
    "N = 10 fog nodes),",
    "rep and seed - which of the five repetitions this is, and the "
    "random seed that generated the instance,",
    "n_variables - total variables in the model: n*m assignment "
    "variables, n*m execution-fraction variables, and n completion-time "
    "variables,",
    "n_constraints - rows in the constraint matrix: one assignment "
    "equation per task, n*m fraction links, and one start-time row per "
    "(task, predecessor, node) triple,",
    "solver_status - 'optimal' (HiGHS proved it), 'time limit' (stopped "
    "at the limit with the best answer found so far), or 'skipped (cap)' "
    "for instances above the 800 assignment-variable limit,",
    "runtime_s - wall-clock seconds for the solve,",
    "optimality_gap - the relative gap the solver still had open when it "
    "stopped; zero means the score was certified optimal,",
    "time_limit_s and rel_gap - the solve settings. Runs that were "
    "certified in task-01 used 300 s and 1e-6; the rest used the "
    "first-pass 60 s and 1e-4.",
]:
    doc.add_paragraph(s, style="List Bullet")

n_runs = len(df)
n_opt = (df.solver_status == "optimal").sum()
n_tl = (df.solver_status == "time limit").sum()
n_skip = (df.solver_status == "skipped (cap)").sum()

doc.add_heading("What the numbers show", 1)
doc.add_paragraph(
    f"The file has {n_runs} rows - one per MILP run reported in the "
    f"paper: {n_opt} finished with a proven-optimal status, {n_tl} "
    f"stopped at the time limit, and {n_skip} were skipped by the "
    "assignment-variable cap (the 500- and 1000-task datasets, where the "
    "model would have thousands of binary variables).")
doc.add_paragraph(
    "The time-limit rows are the interesting ones. They cluster where "
    "the model is genuinely hard: at alpha = 1.0 (the tightest deadline) "
    "and at the degraded bandwidths, where the deadline stays fixed while "
    "transfers get expensive. On those runs the solver was still "
    "improving when the clock stopped - which is exactly the mechanism "
    "behind the task-01 finding that a heuristic occasionally appeared "
    "above the MILP.")

doc.add_heading("A few example rows", 1)
show = df[(df.solver_status != "optimal")].head(8)
t = doc.add_table(rows=1, cols=7)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["experiment", "param", "rep", "n_vars",
                       "n_cons", "status", "gap"]):
    t.rows[0].cells[i].text = h
for _, r in show.iterrows():
    for i, v in enumerate([r.experiment, str(r.param), str(r.rep),
                           str(r.n_variables), str(r.n_constraints),
                           r.solver_status,
                           "" if pd.isna(r.optimality_gap)
                           else f"{r.optimality_gap:.4f}"]):
        t.add_row().cells[i].text = v

doc.add_paragraph(
    "The full table - every experiment, setting, and repetition - is in "
    "the CSV. It was produced by revisions/milp_solver_report.py, which "
    "rebuilds each instance with the same generator arguments and seed "
    "the experiments used, so every number can be regenerated.")

doc.save(os.path.join(HERE, "MILP_solver_details.docx"))
print("wrote revisions/MILP_solver_details.docx")
